"""Build readable DOCX from the evidence-backed Markdown, without altering homework 3."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

root = Path(__file__).resolve().parent
source = root / "report" / "lab4_report.md"
target = root / "report" / "lab4_report_verified.docx"
rows = source.read_text(encoding="utf-8").splitlines()
doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2)
sec.bottom_margin = Cm(2)
sec.left_margin = Cm(2.6)
sec.right_margin = Cm(1.9)
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
for key, font_size in (("Title", 16), ("Heading 1", 13), ("Heading 2", 12)):
    style = doc.styles[key]
    style.font.name = "Times New Roman"
    style.font.size = Pt(font_size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(28, 47, 75)


def rich(paragraph, line):
    pieces = re.split(r"(\*\*.*?\*\*|\x60[^\x60]+\x60)", line)
    for piece in pieces:
        if piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith(chr(96)) and piece.endswith(chr(96)):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif piece:
            paragraph.add_run(piece)


i = 0
while i < len(rows):
    line = rows[i].strip()
    if not line or line == "---":
        i += 1
        continue
    if line.startswith("# "):
        p = doc.add_paragraph(style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rich(p, line[2:])
    elif line.startswith("## "):
        doc.add_heading(line[3:], level=1)
    elif line.startswith("### "):
        doc.add_heading(line[4:], level=2)
    elif line.startswith(chr(96) * 3):
        i += 1
        while i < len(rows) and not rows[i].strip().startswith(chr(96) * 3):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(rows[i])
            run.font.name = "Consolas"
            run.font.size = Pt(7.3)
            i += 1
    elif line.startswith("|"):
        group = []
        while i < len(rows) and rows[i].strip().startswith("|"):
            parts = [c.strip() for c in rows[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", x or "") for x in parts):
                group.append(parts)
            i += 1
        if group:
            cols = max(len(row) for row in group)
            table = doc.add_table(rows=len(group), cols=cols)
            table.style = "Table Grid"
            for index, items in enumerate(group):
                for j, item in enumerate(items):
                    cell = table.cell(index, j)
                    cell.text = item.replace("**", "")
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(0)
                        for run in p.runs:
                            run.font.size = Pt(9)
                            if index == 0:
                                run.bold = True
        continue
    elif line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        rich(p, line[2:])
    else:
        p = doc.add_paragraph()
        rich(p, line)
    i += 1

footer = sec.footer.paragraphs[0]
footer.text = "Домашняя работа № 4 · Apache ZooKeeper · 2026"
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in footer.runs:
    r.font.size = Pt(9)

doc.core_properties.title = "ДЗ 4 — Отказоустойчивость средствами Apache ZooKeeper"
doc.core_properties.subject = "Лабораторная работа, вариант 11, реальные доказательства"
doc.save(target)
print("DOCX_CREATED", target, target.stat().st_size)
